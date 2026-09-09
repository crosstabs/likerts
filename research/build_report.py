from pathlib import Path
import re, json
from urllib.parse import urlparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).parent
source=(root/'report-source.md').read_text()
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(6)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
for element in doc.styles.element.iter():
 for child in list(element):
  if child.tag==qn('w:pBdr'):element.remove(child)

def rich(p,text):
 pos=0
 for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
  p.add_run(text[pos:m.start()])
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m.group(2),RT.HYPERLINK,is_external=True))
  r=OxmlElement('w:r');props=OxmlElement('w:rPr')
  color=OxmlElement('w:color');color.set(qn('w:val'),'175676');props.append(color)
  u=OxmlElement('w:u');u.set(qn('w:val'),'single');props.append(u)
  r.append(props);t=OxmlElement('w:t');t.text=m.group(1);r.append(t);h.append(r);p._p.append(h)
  pos=m.end()
 p.add_run(text[pos:])

def table(lines):
 rows=[[x.strip() for x in l.strip('|').split('|')] for l in lines]
 rows=[r for r in rows if not all(re.fullmatch(r'[- :]+',x) for x in r)]
 t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
 widths=([1.28,1.64,2.34,1.61] if len(rows[0])==4 else [1.48,5.39])
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,txt in enumerate(row):
   cells[j].width=Inches(widths[j]);p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5);rich(p,txt)
   for run in p.runs:run.font.size=Pt(9.5)
   tcPr=cells[j]._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcPr.append(b)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EEF2');tcPr.append(shade)
    for run in p.runs:run.bold=True
  trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  if i==0:trPr.append(OxmlElement('w:tblHeader'))
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=source.splitlines();i=0
while i<len(lines):
 line=lines[i]
 if line.startswith('|'):
  batch=[]
  while i<len(lines) and lines[i].startswith('|'):batch.append(lines[i]);i+=1
  table(batch);continue
 if line=='<!-- PAGE -->':doc.add_page_break()
 elif line.startswith('# '):rich(doc.add_paragraph(style='Title'),line[2:])
 elif line.startswith('## '):rich(doc.add_paragraph(style='Heading 1'),line[3:])
 elif line.startswith('### '):rich(doc.add_paragraph(style='Heading 2'),line[4:])
 elif line.strip():rich(doc.add_paragraph(),line)
 i+=1
doc.core_properties.title='Open source forms for Likerts'
doc.core_properties.subject='Product foundation evaluation'
doc.core_properties.author='Likerts research'
out=root/'Likerts-open-source-forms-evaluation.docx';doc.save(out)
ledger=[]
for paragraph in source.split('\n\n'):
 for label,url in re.findall(r'\[([^\]]+)\]\(([^)]+)\)',paragraph):
  ledger.append({'claim_context':paragraph,'source_title':label,'publisher':urlparse(url).netloc,'url':url,'publication_or_update_date':None,'accessed':'2026-09-06','access_notes':'First-party documentation or repository reviewed; undated unless stated in report.'})
(root/'claim-source-ledger.json').write_text(json.dumps(ledger,indent=2))
print(out);print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'source references',len(ledger))
